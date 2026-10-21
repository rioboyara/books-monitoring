"""Часть 1. Нормализация строк, дубли и сопоставление с книгами.

Запуск: python normalize.py input/<таблица>.xlsx out/result.xlsx [out/pages.json]
pages.json (результат checker.py) необязателен: если он есть, в таблицу
добавляется проверка страниц и заполняется лист "Задание 2".
"""
import hashlib
import json
import re
import sys
from difflib import SequenceMatcher
from collections import Counter, defaultdict
from urllib.parse import unquote, urlsplit

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

SOURCE_COLUMNS = ["Правообладатель", "Автор", "Название", "Тип контента", "url", "Сайт"]

TRANSLIT = dict(zip("абвгдеёжзийклмнопрстуфхцчшщъыьэюя",
                    "a b v g d e e zh z i i k l m n o p r s t u f h c ch sh sch _ y _ e yu ya".split()))
# разные сайты транслитерируют по-разному (x/h/kh, j/y/i, gh/zh), сводим к одному виду
LOOSE = [("shh", "s"), ("sch", "s"), ("tch", "c"), ("sh", "s"), ("ch", "c"), ("zh", "z"), ("gh", "z"),
         ("kh", "h"), ("ts", "c"), ("tc", "c"), ("x", "h"), ("j", "i"), ("y", "i"), ("w", "v"), ("q", "k")]
AUDIO_URL = re.compile(r"mp3|audio|track|slushat|zvuk", re.I)
NOT_BOOK_URL = re.compile(r"/tag/|section_id=|sbornik|/avtor|/author", re.I)


def clean(value):
    """Значение для показа: без лишних пробелов и мусора по краям (;100 ГОТОВЫХ, Цветок Жизни;)."""
    return re.sub(r"\s+", " ", str(value or "")).strip(" ;*\t")


def key(text):
    """Ключ для сравнения: только буквы и цифры, без пробелов. Ценадоверия == Цена доверия."""
    return re.sub(r"[\W_]+", "", clean(text).lower().replace("ё", "е"))


def loose(text):
    """Латиница в одном написании, чтобы сравнивать название с адресом страницы."""
    s = "".join(TRANSLIT.get(ch, ch) for ch in unquote(str(text)).lower()).replace("_", "")
    for a, b in LOOSE:
        s = s.replace(a, b)
    return re.sub(r"(.)\1+", r"\1", re.sub(r"[^a-z0-9 ]+", " ", s)).strip()


URL_SERVICE_WORDS = {loose(w) for w in "book books html php threads viewtopic story index read reader".split()}


def split_author(raw):
    """Из '(JustHope) Егорова Надежда;' получается ('Егорова Надежда', 'JustHope')."""
    raw = clean(raw)
    aliases = re.findall(r"\(([^)]*)\)", raw)
    name = clean(re.sub(r"\([^)]*\)", " ", raw))
    return name, ", ".join(aliases)


def author_key(name):
    """Слова имени по алфавиту, с сохранением инициалов."""
    words = re.findall(r"\w+", name.lower().replace("ё", "е"))
    return " ".join(sorted(words))


def book_id(author, title):
    """Стабильный id: зависит только от автора и названия, поэтому не меняется при новых строках."""
    raw = author_key(split_author(author)[0]) + "|" + key(title)
    return "B" + hashlib.md5(raw.encode()).hexdigest()[:6]


def author_matches(expected, actual):
    """Одинаковый алфавит требует точного имени; допуск нужен только для транслитерации."""
    expected_words = re.findall(r"\w+", clean(expected).lower().replace("ё", "е"))
    actual_words = set(re.findall(r"\w+", clean(actual).lower().replace("ё", "е")))
    for word in expected_words:
        if word in actual_words:
            continue
        cross_script = [p for p in actual_words if bool(re.search(r"[а-я]", p)) != bool(re.search(r"[а-я]", word))]
        if not any(loose(word)[:1] == loose(p)[:1] and SequenceMatcher(None, loose(word), loose(p)).ratio() >= .8
                   for p in cross_script):
            return False
    return bool(expected_words)


def url_key(url):
    """Одна и та же страница: без зеркал s4./www., без index.php, хвостового слеша и меток."""
    p = urlsplit(clean(url))
    host = re.sub(r"^(www|s\d+)\.", "", (p.hostname or "").lower())
    path = unquote(p.path).replace("/index.php", "").rstrip("/")
    query = "&".join(q for q in p.query.split("&") if q and not q.startswith(("utm_", "dsource")))
    return host + path + ("?" + query if query else "")


def title_vs_url(title, url):
    """Доля слов названия, которые видны в адресе страницы (None, если в адресе только номер), и пропавшие номера томов."""
    slug_words = [w for w in loose(urlsplit(url).path + " " + urlsplit(url).query).split()
                  if len(w) > 2 and not w.isdigit() and w not in URL_SERVICE_WORDS]
    if not slug_words:
        return None, []
    slug = "".join(slug_words)
    words = [w for w in loose(title).split() if len(w) > 2]
    share = sum(w in slug for w in words) / len(words) if words else 1.0
    slug_numbers = {int(n) for n in re.findall(r"\d+", unquote(url))}
    lost = [n for n in re.findall(r"\d+", title) if int(n) not in slug_numbers]
    return share, lost


def load_rows(src, sheet="Задание 1"):
    """src: путь, файл или уже открытая книга openpyxl."""
    ws = (src if isinstance(src, openpyxl.Workbook) else openpyxl.load_workbook(src))[sheet]
    rows = []
    for n, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if any(values[:6]):
            original = dict(zip(SOURCE_COLUMNS, (*values[:6], *([None] * max(0, 6 - len(values))))))
            rows.append({"row": n, **original, "original": original, "extra": values[6:]})
    return rows


def process(rows):
    """Добавляет к строкам нормализованные поля, id книги, дубли и пометки."""
    for r in rows:
        r["author"], r["alias"] = split_author(r["Автор"])
        r["title"] = clean(r["Название"])
        r["url"] = clean(r["url"])
        r["url_key"] = url_key(r["url"])
        r["site"] = clean(r["Сайт"]) or urlsplit(r["url"]).hostname
        r["book"] = book_id(r["Автор"], r["Название"])
        r["notes"], r["doubts"] = [], []

    by_book = defaultdict(list)
    for r in rows:
        by_book[r["book"]].append(r)

    first_seen = {}
    for r in rows:
        k = (r["url_key"], r["book"])
        r["dup_of"] = first_seen.setdefault(k, r["row"])
        if r["dup_of"] != r["row"]:
            same = clean(r["url"]) == clean(next(x["url"] for x in rows if x["row"] == r["dup_of"]))
            r["notes"].append(f"дубль строки {r['dup_of']}" + ("" if same else " (тот же адрес на зеркале)"))

    books_on_page = defaultdict(set)
    for r in rows:
        books_on_page[r["url_key"]].add(r["book"])

    authors_by_title = defaultdict(set)
    for r in rows:
        authors_by_title[key(r["title"])].add(r["author"])

    author_keys = {author_key(r["author"]) for r in rows}
    spellings = defaultdict(set)
    for r in rows:
        spellings[author_key(r["author"])].add(r["author"])

    for r in rows:
        if any(clean(r[c]) != str(r[c] or "").strip() for c in ("Автор", "Название")):
            r["notes"].append("убраны лишние символы")
        if r["alias"]:
            r["notes"].append(f"псевдоним в скобках: {r['alias']}")
        if len(spellings[author_key(r["author"])]) > 1:
            r["notes"].append("автор записан по-разному: " + " / ".join(sorted(spellings[author_key(r["author"])])))
        if re.search(r"[a-z]", r["author"], re.I) and re.search(r"[а-я]", r["author"], re.I):
            r["doubts"].append("в поле автора два написания, латиницей и кириллицей")
        ak = set(author_key(r["author"]).split())
        bigger = [a for a in author_keys if ak < set(a.split())]
        if bigger:
            r["doubts"].append("возможно, тот же автор, что " + ", ".join(sorted(min(spellings[a]) for a in bigger)))
        others = authors_by_title[key(r["title"])] - {r["author"]}
        if others:
            r["doubts"].append("это же название есть у других авторов: " + ", ".join(sorted(others)))
        if len(books_on_page[r["url_key"]]) > 1:
            r["doubts"].append("одна страница указана для разных книг")
        share, lost = title_vs_url(r["title"], r["url"])
        r["url_match"] = "" if share is None else round(share * 100)
        if share is None:
            r["notes"].append("в адресе нет названия, сверить можно только по самой странице")
        elif share < 0.5:
            r["doubts"].append(f"адрес страницы не похож на название (совпало {r['url_match']}% слов)")
        if lost and share and share >= 0.5:
            r["doubts"].append("номер тома или части в адресе другой: в названии " + ", ".join(lost))
        if NOT_BOOK_URL.search(r["url"]):
            r["doubts"].append("похоже на страницу тега, раздела или сборника, а не одной книги")
        if AUDIO_URL.search(r["url"]) and str(r["Тип контента"]).lower() == "ebook":
            r["doubts"].append("тип ebook, а ссылка похожа на аудио")
    return rows


def catalog(rows):
    """Список книг: id, автор, название и все строки, где книга встречается."""
    books = {}
    for r in rows:
        b = books.setdefault(r["book"], {"id": r["book"], "rows": [], "titles": Counter(), "authors": Counter(),
                                          "pages": set(), "sites": set()})
        b["rows"].append(r["row"])
        b["titles"][r["title"]] += 1
        b["authors"][r["author"]] += 1
        b["pages"].add(r["url_key"])
        b["sites"].add(r["site"])
    for b in books.values():
        b["title"] = b["titles"].most_common(1)[0][0]
        b["author"] = b["authors"].most_common(1)[0][0]
    return books


# какие данные закрыли бы вопрос, по тексту пометки
NEEDS = [
    ("адрес страницы не похож", "ID книги у правообладателя (артикул или ISBN) для каждой ссылки"),
    ("номер тома", "серия и номер тома отдельными полями"),
    ("есть у других авторов", "ID книги или ISBN, чтобы различать одноимённые книги"),
    ("тот же автор", "список псевдонимов и вариантов имени автора"),
    ("два написания", "основное имя автора и варианты написания"),
    ("тега, раздела или сборника", "ссылка на страницу конкретной книги или состав сборника"),
    ("похожа на аудио", "на какие форматы есть права: текст, аудио"),
    ("другая книга", "подтверждение правообладателя, к какой книге относится ссылка"),
    ("автора из каталога", "варианты имени автора и псевдонимы"),
    ("в списке на странице", "ссылка на страницу конкретной книги"),
    ("Книг из каталога на странице нет", "проверить, та ли ссылка указана для книги"),
    ("совпадение с каталогом не найдено", "проверить, та ли ссылка указана для книги"),
    ("Страницы нет", "не нужны: страница удалена, ссылку можно закрыть"),
    ("покупку в Литрес", "не нужны: страница партнёрская"),
    ("скачать, читать или слушать нельзя", "не нужны: размещения нет"),
]

def page_state(p):
    return ", ".join(str(x) for x in (p.get("http_status"), p.get("page_state")) if x) or "нет ответа"


def ai_summary(p):
    ai = p.get("ai")
    if not ai:
        return ""
    return "ИИ: " + ("решение уточнено, цитаты проверены" if ai.get("accepted") else ai.get("note", "нужна ручная проверка"))


def extracted(p):
    """Что прототип достал со страницы: заголовок, автор, тип файла."""
    return "; ".join(filter(None, [
        f"заголовок: {p['heading'][:150]}" if p.get("heading") else "",
        f"автор на странице: {p['page_author']}" if p.get("page_author") else "",
        f"загрузка: {p['extraction_method']}" if p.get("extraction_method") else "",
        f"тип: {p['content_type']}" if p.get("content_type") and "html" not in p["content_type"] else "", ai_summary(p)]))


def needs(text):
    return "; ".join(dict.fromkeys(n for k, n in NEEDS if k in text))


def write_xlsx(rows, books, path, part2=None, pages=None, advice=None):
    pages = pages or {}
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    head = Font(bold=True)
    fill = PatternFill("solid", fgColor="FFF2CC")

    def sheet(title, header, data, widths):
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font = head
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for line in data:
            ws.append(line)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[ws.max_row].height = 58
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    def page_cols(url):
        p = pages.get(url)
        if not p:
            return ["", "", ""]
        found = p["match"]["book_id"] if p.get("match") else ""
        return [p["status"], p["reason"], found]

    ws = sheet("Строки", [
        "Строка в исходнике", *SOURCE_COLUMNS, "Автор (чистый)", "Псевдоним", "Название (чистое)",
        "Ключ страницы", "ID книги", "Дубль строки", "Слов названия в адресе, %", "Пометки", "Спорно",
        "Проверка страницы: статус", "Проверка страницы: причина", "Книга на странице (ID)"],
        [[r["row"], *[r.get("original", r)[c] for c in SOURCE_COLUMNS], r["author"], r["alias"], r["title"], r["url_key"],
          r["book"], "" if r["dup_of"] == r["row"] else r["dup_of"], r["url_match"], "; ".join(r["notes"]),
          "; ".join(r["doubts"]), *page_cols(r["url"])] for r in rows],
        [9, 12, 20, 30, 8, 30, 16, 20, 12, 30, 35, 10, 9, 10, 30, 40, 9, 45, 11])
    for line in ws.iter_rows(min_row=2):
        if line[15].value:
            for c in line:
                c.fill = fill

    sheet("Книги", ["ID книги", "Автор", "Название", "Строк", "Разных страниц", "Строки в исходнике",
                    "Варианты записи автора", "Варианты названия", "Сайты"],
          [[b["id"], b["author"], b["title"], len(b["rows"]), len(b["pages"]), ", ".join(map(str, b["rows"])),
            " / ".join(b["authors"]) if len(b["authors"]) > 1 else "",
            " / ".join(b["titles"]) if len(b["titles"]) > 1 else "", ", ".join(sorted(b["sites"]))]
           for b in sorted(books.values(), key=lambda b: (b["author"], b["title"]))],
          [10, 25, 45, 7, 9, 18, 30, 30, 40])

    sheet("Дубли", ["Строка", "Дубль строки", "Автор", "Название", "url", "Почему дубль"],
          [[r["row"], r["dup_of"], r["author"], r["title"], r["url"], "; ".join(n for n in r["notes"] if "дубль" in n)]
           for r in rows if r["dup_of"] != r["row"]],
          [8, 10, 22, 40, 60, 35])

    doubtful = []
    for r in rows:
        p = pages.get(r["url"])
        page_doubt = ""
        # технические сбои (бот-защита, таймаут) не спорность данных, они видны в статусе на листе "Строки"
        if p and p.get("match") and p["match"]["book_id"] != r["book"]:
            page_doubt = f"на странице другая книга: {p['match']['author']}. {p['match']['title']}"
        elif p and p["page_state"] in ("ok", "страницы нет") and p["status"] != 1:
            page_doubt = f"страница не подтвердила книгу (статус {p['status']}): {p['reason']}"
        if r["doubts"] or page_doubt:
            data_doubt = "; ".join(r["doubts"])
            doubtful.append([r["row"], r["author"], r["title"], r["url"], data_doubt, page_doubt,
                             needs(data_doubt + " " + page_doubt)])
    sheet("Спорные", ["Строка", "Автор", "Название", "url", "Что не так по данным", "Что показала страница",
                      "Каких данных не хватает"],
          doubtful, [8, 22, 40, 50, 55, 55, 45])

    if part2:
        sheet("Задание 2", ["Строка в исходнике", *SOURCE_COLUMNS, "Пират-1/не пират-0 (статус 0/1/2)", "Признаки",
                            "Статус страницы", "Извлечённые данные", "Совпадение с каталогом", "Причина решения"],
              [[r["row"], *[r.get("original", r)[c] for c in SOURCE_COLUMNS], *part2_cols(pages.get(r["url"]))] for r in part2],
              [9, 12, 20, 30, 8, 40, 18, 12, 40, 16, 50, 30, 60])
    if advice is not None:
        sheet("Рекомендации ИИ", ["Строка в исходнике", "Название", "Почему строка спорная", "Что проверить",
                                "Каких данных не хватает", "Источник рекомендации"], advice, [10, 30, 50, 55, 45, 30])
    wb.save(path)


def part2_cols(p):
    if not p:
        return ["", "", "не проверено", "", "", ""]
    m = p.get("match")
    return [p["status"], "; ".join(p["signs"]), page_state(p), extracted(p),
            f"{m['author']}. {m['title']}" if m else "нет", p["reason"]]


def build(src, dst, pages=None, advise=None):
    """Собирает таблицу результата. src и dst: пути или файлы. pages: результаты checker.py по url."""
    wb = src if isinstance(src, openpyxl.Workbook) else openpyxl.load_workbook(src)
    rows = process(load_rows(wb))
    books = catalog(rows)
    part2 = load_rows(wb, "Задание 2") if "Задание 2" in wb.sheetnames else []
    for r in part2:
        r["url"] = clean(r["url"])
    advice = advise(rows) if advise else None
    write_xlsx(rows, books, dst, part2, pages or {}, advice)
    return rows, books


def main(src, dst, pages_path=None):
    pages = {}
    if pages_path:
        with open(pages_path, encoding="utf-8") as f:
            pages = {p["url"]: p for p in json.load(f)}
    rows, books = build(src, dst, pages)
    dups = sum(r["dup_of"] != r["row"] for r in rows)
    print(f"строк {len(rows)}, книг {len(books)}, дублей {dups}, спорных {sum(bool(r['doubts']) for r in rows)}")


if __name__ == "__main__":
    main(*sys.argv[1:4])
