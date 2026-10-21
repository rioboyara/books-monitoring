"""Проверка обработки файлов: python test_uploads.py (нужен httpx)."""
import io
import zipfile
from unittest.mock import patch
from pathlib import Path
import openpyxl
from fastapi.testclient import TestClient
from app import app, MAX_UNPACKED_BYTES
from normalize import load_rows, process, catalog


def xlsx_bytes(wb):
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def test_uploads():
    client = TestClient(app)
    src = next(Path('input').glob('*.xlsx'))
    good = client.post('/api/table', content=src.read_bytes())
    assert good.status_code == 200, good.text
    wb = openpyxl.load_workbook(io.BytesIO(good.content))
    assert wb['Строки'].max_row == 151
    assert client.post('/api/table', content=b'not xlsx').status_code == 400
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, 'w', zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('big.xml', b'0' * (MAX_UNPACKED_BYTES + 1))
    assert client.post('/api/table', content=bomb.getvalue()).status_code == 400
    source = openpyxl.load_workbook(src)
    source['Задание 2'].cell(3001, 1, 'too many')
    assert client.post('/api/table', content=xlsx_bytes(source)).status_code == 400
    source = openpyxl.load_workbook(src)
    source['Задание 1']['E2'] = '  https://example.org/book  '
    source['Задание 1']['C2'] = '=1+1'
    result = client.post('/api/table', content=xlsx_bytes(source))
    assert result.status_code == 200
    output = openpyxl.load_workbook(io.BytesIO(result.content))['Строки']
    assert output['F2'].value == '  https://example.org/book  '
    assert output['D2'].value == '=1+1' and output['D2'].data_type == 's'
    with patch('app.suggest_rows', return_value=[[2, 'Книга', 'Проверить автора', 'ID книги']]):
        assisted = client.post('/api/table', content=src.read_bytes(), headers={'X-AI-Review': '1'})
        assert assisted.status_code == 200
        assisted_wb = openpyxl.load_workbook(io.BytesIO(assisted.content))
        assert assisted_wb['Рекомендации ИИ']['A2'].value == 2
        assert assisted_wb['Строки']['D2'].value == wb['Строки']['D2'].value
    rows = process(load_rows(src))
    print('source counts', len(rows), len(catalog(rows)), sum(r['dup_of'] != r['row'] for r in rows))


if __name__ == '__main__':
    test_uploads()
    print('upload checks passed')
