"""Веб-обёртка: записка на главной, проверка ссылок на /check, обработка своей таблицы на /table."""
import io
import json
import zipfile
from pathlib import Path
from threading import BoundedSemaphore

import openpyxl
from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles
from starlette.concurrency import run_in_threadpool
from pydantic import BaseModel

from checker import check_many, load_catalog
from normalize import build

HERE = Path(__file__).parent
MAX_URLS = 30
MAX_TABLE_BYTES = 5_000_000
MAX_TABLE_ROWS = 3000
MAX_UNPACKED_BYTES = 50_000_000
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
busy = BoundedSemaphore(2)

app = FastAPI(docs_url=None, redoc_url=None)
catalog = load_catalog(next((HERE / "input").glob("*.xlsx")))
app.mount("/files", StaticFiles(directory=HERE / "out"), name="files")
app.mount("/assets", StaticFiles(directory=HERE / "web"), name="assets")


class Urls(BaseModel):
    urls: list[str]


@app.get("/")
def note():
    return FileResponse(HERE / "web" / "index.html")


@app.get("/check")
def check_page():
    return FileResponse(HERE / "web" / "check.html")


@app.get("/table")
def table_page():
    return FileResponse(HERE / "web" / "table.html")


@app.post("/api/check")
def api_check(data: Urls):
    urls = list(dict.fromkeys(u.strip() for u in data.urls if u.strip()))
    if not urls or len(urls) > MAX_URLS or not all(u.startswith(("http://", "https://")) for u in urls):
        raise HTTPException(400, f"Нужен список от 1 до {MAX_URLS} адресов, каждый с http:// или https://")
    # доступ к сайту закрыт логином в nginx, здесь только ограничение одновременных проверок
    if not busy.acquire(blocking=False):
        raise HTTPException(429, "Сервер занят другой проверкой, попробуйте через минуту")
    try:
        return check_many(urls, catalog)
    finally:
        busy.release()


@app.post("/api/table")
async def api_table(request: Request):
    """Своя версия исходной таблицы, например с новыми строками: возвращает обработанную таблицу.
    Для адресов, которые уже проверялись, подставляется результат проверки страницы."""
    body = await request.body()
    if not body or len(body) > MAX_TABLE_BYTES:
        raise HTTPException(400, "Нужен файл xlsx до 5 МБ")
    # разбор xlsx занимает процессор, поэтому в отдельном потоке и под тем же ограничением, что и проверка ссылок
    return await run_in_threadpool(make_table, body)


def make_table(body):
    bad_file = HTTPException(400, 'Не удалось прочитать файл. Нужен xlsx с листом "Задание 1" в формате исходной таблицы')
    try:
        # xlsx это zip: маленький файл может распаковаться в гигабайты, поэтому смотрим размер до разбора
        unpacked = sum(i.file_size for i in zipfile.ZipFile(io.BytesIO(body)).infolist())
    except zipfile.BadZipFile:
        raise bad_file
    if unpacked > MAX_UNPACKED_BYTES:
        raise HTTPException(400, "Файл слишком большой после распаковки")
    if not busy.acquire(blocking=False):
        raise HTTPException(429, "Сервер занят другой задачей, попробуйте через минуту")
    try:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(body))
            sheets = [wb[name] for name in ("Задание 1", "Задание 2") if name in wb.sheetnames]
        except Exception:  # чужой файл может быть любым, отвечаем понятной ошибкой
            raise bad_file
        if "Задание 1" not in wb.sheetnames:
            raise bad_file
        if any(ws.max_row > MAX_TABLE_ROWS or ws.max_column > 30 for ws in sheets):
            raise HTTPException(400, f"Не больше {MAX_TABLE_ROWS} строк и 30 колонок на листе")
        pages_file = HERE / "out" / "pages.json"
        pages = {p["url"]: p for p in json.loads(pages_file.read_text(encoding="utf-8"))} if pages_file.exists() else {}
        out = io.BytesIO()
        try:
            build(wb, out, pages)
        except Exception:
            raise bad_file
    finally:
        busy.release()
    return Response(out.getvalue(), media_type=XLSX, headers={"Content-Disposition": 'attachment; filename="result.xlsx"'})
